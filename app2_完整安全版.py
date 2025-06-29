import streamlit as st
import pandas as pd
import time
from openai import OpenAI
import json
from io import BytesIO
from docx import Document
from fpdf import FPDF
from openpyxl import Workbook
from sklearn.model_selection import train_test_split
from sklearn.ensemble import RandomForestClassifier
from sklearn.preprocessing import LabelEncoder
import matplotlib.pyplot as plt
import seaborn as sns
import re
import os
from dotenv import load_dotenv
from tempfile import NamedTemporaryFile

# ✅ 頁面設定
st.set_page_config(page_title="GPT AI 全功能極速助手", layout="wide", page_icon="⚡")
st.markdown("<h1 style='text-align: center; color: #2C3E50;'>⚡ GPT AI 全功能極速助手</h1>", unsafe_allow_html=True)

# ✅ OpenAI API
load_dotenv()
api_key = os.getenv("OPENAI_API_KEY")
client = OpenAI(api_key=api_key)

# ✅ 特休資訊
annual_leave_info = "📅 特休依《勞基法》第38條：滿6個月3天、滿1年7天、滿2年10天…最高30天"

# ✅ 功能提示詞
prompt_map = {
    "履歷表產生": "請幫我根據以下資訊，產生一份正式中文履歷表：",
    "專案計畫書": "請撰寫一份完整的專案計畫書，內容包含目的、目標、執行步驟、時間表與預期成果：",
    "合約草稿": "請撰寫以下需求對應的合約草稿：",
    "出勤紀錄表": "請幫我建立一份包含員工出勤紀錄的表格格式：",
    "資料分析報表": "請根據以下資料建立摘要、趨勢與建議：",
    "函式 (google sheet+excel函數公式)": "請根據下列需求，僅推薦最相關的 Google Sheets 或 Excel 函數，並提供簡要範例與用途：",
    "函式+解說 (google sheet+excel函數公式)": "請根據下列問題，只推薦必要的 Google Sheets / Excel 函數，並說明用途與範例：",
    "教學生成(google sheet+excel函數公式)": "請根據下列描述，產生教學與實例，只包含符合需求的 Google Sheets / Excel 函數：",
    "法律諮詢": "請針對以下法律問題提供意見，並標註法規依據：",
    "特休公式選項": annual_leave_info + "\n請產生 Google Sheets / Excel 特休計算公式",
    "特休公式+解說選項": annual_leave_info + "\n請產生公式並說明用途與欄位設定",
    "特休教學生成選項": annual_leave_info + "\n請產生教學與使用流程，並提供公式",
    "寫 Python 程式": "請**只用 Python**寫以下需求的程式，並簡要說明教學。",
    "寫 Apps Script 程式": "請**只用 Google Apps Script**寫以下需求的程式，並簡要說明教學。",
    "翻譯選項/英文/韓文/日文/法文/": "請將以下文字翻譯為英文、韓文、日文與法文：",
    "產生英文報告": "請根據以下資料撰寫一份正式英文報告：",
    "產生韓文報告": "請根據以下資料撰寫一份正式韓文報告：",
    "產生日文報告": "請根據以下資料撰寫一份正式日文報告：",
    "自動生成Ptt文案": "請幫我根據以下主題，自動生成一篇風格類似 Ptt 鄉民的推文文案："
}

def clean_response(text):
    text = re.sub(r"[、‧•．●【】「」『』（）()]", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text

def save_as_word(content):
    doc = Document()
    for line in content.split('\n'):
        doc.add_paragraph(line)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def save_as_pdf(content):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for line in content.split('\n'):
        pdf.multi_cell(0, 10, line)
    buffer = BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer

def save_as_excel(content):
    wb = Workbook()
    ws = wb.active
    for i, line in enumerate(content.split('\n')):
        ws.cell(row=i + 1, column=1, value=line)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

# ✅ AI 缺勤預測
st.header("📂 AI 缺勤預測")
uploaded_file = st.file_uploader("上傳包含欄位的 Excel", type=["xlsx"])
if uploaded_file:
    df = pd.read_excel(uploaded_file)
    st.dataframe(df)
    required_cols = ["部門", "班別", "出勤天數", "請假天數", "遲到次數", "是否異常"]
    if all(col in df.columns for col in required_cols):
        data = df.copy()
        le_dict = {}
        for col in ["部門", "班別"]:
            le = LabelEncoder()
            data[col] = le.fit_transform(data[col])
            le_dict[col] = le
        X = data[["部門", "班別", "出勤天數", "請假天數", "遲到次數"]]
        y = data["是否異常"]
        X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2)
        model = RandomForestClassifier(n_estimators=100)
        model.fit(X_train, y_train)
        df["風險預測"] = model.predict(X)
        df["風險預測"] = df["風險預測"].map({0: "正常", 1: "高風險"})
        st.success(f"✅ 模型準確率：{model.score(X_test, y_test) * 100:.2f}%")
        st.dataframe(df)
        st.download_button("📥 下載預測結果", df.to_csv(index=False).encode("utf-8-sig"), file_name="風險預測.csv")
        st.subheader("📊 風險分佈圖")
        fig, ax = plt.subplots()
        sns.countplot(data=df, x="部門", hue="風險預測", ax=ax)
        st.pyplot(fig)
    else:
        st.warning("⚠️ 請包含欄位：" + ", ".join(required_cols))

# ✅ GPT 生成功能區
st.header("🧠 GPT 極速生成功能")

if 'user_input' not in st.session_state:
    st.session_state.user_input = ""

if 'result' not in st.session_state:
    st.session_state.result = None

def update_user_input():
    st.session_state.user_input = ""

feature = st.selectbox("🎯 選擇功能", list(prompt_map.keys()), key="feature", on_change=update_user_input)
model = st.selectbox("💡 選擇 GPT 模型", ["gpt-4o", "gpt-4", "gpt-4.1-nano", "gpt-3.5-turbo"])
user_input = st.text_area("✍️ 輸入內容（請描述你要做的事）", height=200, key="user_input")

if st.button("⚡ 開始即時產生"):
    if not user_input.strip():
        st.warning("⚠️ 請輸入內容")
    else:
        full_prompt = prompt_map[feature] + "\n" + user_input.strip()
        st.info("📝 以下為輸出結果：")
        with st.spinner("AI 正在生成中...請稍候"):
            response = client.chat.completions.create(
                model=model,
                messages=[{"role": "user", "content": full_prompt}],
                temperature=0.5
            )
            result = response.choices[0].message.content
            st.session_state.result = result
            try:
                st.experimental_rerun()
            except AttributeError:
                pass

# ✅ 顯示結果與播放語音（不會因為按鈕被清除）
if st.session_state.result:
    st.markdown(
        f"<div style='background-color:#F4F6F6;padding:10px;border-radius:5px;'>"
        f"<pre style='white-space: pre-wrap;'>{st.session_state.result}</pre></div>",
        unsafe_allow_html=True
    )

    # 🔊 語音語言下拉選單
    lang_options = {
        "中文": "zh",
        "英文": "en",
        "日文": "ja",
        "韓文": "ko"
    }
    lang_name = st.selectbox("選擇語音朗讀語言", list(lang_options.keys()), index=0)
    lang = lang_options[lang_name]

    if st.button("🔊 播放語音"):
        speech_response = client.audio.speech.create(
            model="tts-1",
            voice="alloy",
            input=st.session_state.result,
            response_format="mp3",
        )
        with NamedTemporaryFile(delete=False, suffix=".mp3") as tmpfile:
            tmpfile.write(speech_response.read())
            tmpfile.flush()
            st.audio(tmpfile.name, format="audio/mp3")

    # 📄 檔案下載
    if feature in ["履歷表產生", "專案計畫書", "合約草稿"]:
        st.download_button("📄 下載 Word", save_as_word(st.session_state.result), file_name="輸出.docx")
        st.download_button("🧾 下載 PDF", save_as_pdf(st.session_state.result), file_name="輸出.pdf")
    elif feature in ["出勤紀錄表", "資料分析報表"]:
        st.download_button("📊 下載 Excel", save_as_excel(st.session_state.result), file_name="輸出.xlsx")
